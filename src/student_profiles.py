import io
import os
import docx
import pandas as pd
from functools import lru_cache
from PIL import Image, ImageDraw
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Pt, RGBColor


CURRENT_DIRECTORY = os.path.dirname(os.path.realpath(__file__)).replace("\\", "/")
OUTPUT_DIRECTORY = f"{CURRENT_DIRECTORY}/student_profiles"
REGISTER_FILE_NAME = "register_students.csv"
REGISTER_FILE_PATH = f"{CURRENT_DIRECTORY}/{REGISTER_FILE_NAME}"
IMAGES_DIRECTORY = f"{os.path.dirname(CURRENT_DIRECTORY)}/images"
LOGO_FILE_PATH = f"{IMAGES_DIRECTORY}/Logo_white.png"  # in the header
WATERMARK_FILE_PATH = f"{IMAGES_DIRECTORY}/Logo.png"  # in the corner of every page


class Theme:
    def __init__(self, card_color: str, frame_color: str, frame_texture_file_path: str):
        self.card_color = card_color
        self.frame_color = frame_color  # the main color of the texture
        self.frame_texture_file_path = frame_texture_file_path


# Look of the profile, from the "Hello Students" deck: rounded cards in a rounded frame on a dark page.
# The frame is a slide background of the deck, with white brush strokes.
# The student's data is on warm cards in the cold frame, the information from the team is on cold cards in the warm one.
WARM_THEME = Theme(card_color="FFC696", frame_color="97A7FF",
                   frame_texture_file_path=f"{IMAGES_DIRECTORY}/background_cold.png")  # peach card, periwinkle frame
COLD_THEME = Theme(card_color="97A6FF", frame_color="FFC596",
                   frame_texture_file_path=f"{IMAGES_DIRECTORY}/background_warm.png")  # periwinkle card, peach frame
FONT_NAME = "Arial"  # the deck fonts are used only if they are installed on every PC that opens the profiles
PAGE_COLOR = "191643"  # dark navy
PAGE_TEXT_COLOR = "FFFFFF"
SUBTITLE_COLOR = "FFC696"  # peach
ACCENT_COLOR = "F64C10"  # orange, for the arrows and the section numbers
CARD_TEXT_COLOR = "130F33"  # navy, also for the timeline on the left of the cards
# Paragraph styles of the fields in the cards
FIELD_TITLE_STYLE = "Profile Field Title"
FIELD_VALUE_STYLE = "Profile Field Value"
PROFILE_TITLE = "Профил на твоя ученик"
PROFILE_SUBTITLE = "Запознай се с ученика, с когото ще работите заедно в ABLE Mentor"

# Sizes
CONTENT_WIDTH = Cm(17)  # between the page margins
FRAME_WIDTH = Cm(0.45)  # around the card
FRAME_RADIUS = Cm(0.85)
CARD_RADIUS = Cm(0.55)
TIMELINE_WIDTH = Cm(0.9)  # the dotted line is in its middle
TIMELINE_DOT_SIZE = 12  # in 1/8 pt
TIMELINE_RING_SIZE = Cm(0.3)
TEXT_WIDTH = CONTENT_WIDTH - 2 * FRAME_WIDTH - TIMELINE_WIDTH
WATERMARK_WIDTH = Cm(11)
IMAGE_DPI = 300  # of the generated images (the rounded corners and the rings)

# Only students with one of these statuses get a profile (matched case-insensitive)
PROFILE_STATUSES = ["in process", "matched"]

# Column names in the students' csv file (matched case-insensitive, long names by their beginning)
STUDENT_FIRST_NAME = "Име"
STUDENT_LAST_NAME = "Фамилия"
STUDENT_STATUS = "Статус"

# Sections of the profile: (section title, fields). Each field is a row of the section's card: (title, columns).
# The values of several columns are joined on separate lines. Empty rows and sections are skipped.
PROFILE_SECTIONS = [
    ("Основна информация", [
        ("Удобен формат", ["Удобен Формат"]),
        ("В кой град предпочитате", ["В кой град предпочитате"]),
        ("Населено място", ["Населено Място"]),
        ("Възраст", ["Възраст"]),
        ("Ниво на английски език", ["Английски Език Ниво"]),
    ]),
    ("Образование и кариера", [
        ("Образование и кариера", ["Образование И Кариера"]),
        ("Училище", ["Училище"]),
        ("Завършен клас", ["Последно Завършен Клас"]),
        ("Любими предмети и дейности в училище", ["Любими Предмети И Дейности В Училище"]),
        ("Какво ще правя след гимназията", ["Какво Ще Правиш След Гимназията"]),
    ]),
    ("За теб извън училище", [
        ("За теб извън училище", ["За Теб Извън Училище"]),
        ("Интереси извън училище", ["Интереси Извън Училище"]),
        ("Разкажи ни за лично предизвикателство и как си го превъзмогнал/а?",
         ["Разкажете Ни За Ваше Лично Предизвикателство"]),
        ("Кой би бил ти в екип? (подредени по предпочитание)",
         ["Кой Би Бил Ти В Екип", "Избор 2", "Избор 3", "Избор 4", "Избор  5", "Избор 6"]),
    ]),
    ("Участие в програмата", [
        ("По колко часа седмично би отделял/а на проекта?", ["Средно по колко часа"]),
        ("Защо имаш нужда да участваш в програмата?", ["Защо имаш нужда да участваш"]),
        ("Участвал ли си в други сходни програми, завършил ли си ги и какво си взе от тях?",
         ["Участвал ли си в други сходни програми"]),
        ("По какъв проект би работил/а със своя ментор? (подредени по предпочитание)",
         ["По какъв проект", "проект 2", "проект 3", "проект 4"]),
        ("Можеш да добавиш няколко аспекта и да ги подредиш по важност.",
         ["Можеш да добавиш няколко аспекта и да ги подредиш по важност."]),
        ("Кое твърдение за теб по отношение на кариерната и академичната насока е вярно?",
         ["Кое твърдение за теб по отношение на кариерната"]),
        ("Кое твърдение за теб по отношение на участието ти в програмата е вярно?",
         ["Кое твърдение за теб по отношение на участието"]),
        ("Кажи повече за това твърдение", ["кажи повече за това твърдение"]),
    ]),
    ("Сфери и ментор", [
        ("Сфери, които са ти интересни, и/или в които искаш да се развиеш",
         ["Сфери, които са ти интересни", "сфера 1", "повече за сфера 1", "сфера 2", "повече за сфера 2", "сфера 3",
          "повече за сфера 3"]),
        # pandas renames the 2nd of two equal headers to "<header>.1"
        ("Можеш да добавиш няколко аспекта и да ги подредиш по важност.",
         ["Можеш да добавиш няколко аспекта и да ги подредиш по важност..1"]),
        ("Има ли специфична сфера, различна от избраните, в която би искал/а менторът ти да има опит и защо?",
         ["Има ли специфична сфера"]),
        ("Кои аспекти на работата с ментора ще са най-полезни? (подредени по важност)",
         ["Кои аспекти на работата с ментора", "избери аспект 1", "избери аспект 2"]),
        ("Научил/а за ABLE Mentor от?", ["Откъде разбрахте"]),
    ]),
]


def read_register(csv_file_path: str) -> pd.DataFrame:
    register = pd.read_csv(csv_file_path, dtype=str, encoding="utf-8").fillna("")
    register.columns = [column.strip() for column in register.columns]
    return register


def find_column(register: pd.DataFrame, column_name: str) -> str:
    column_name = column_name.casefold()
    for column in register.columns:
        if column.casefold() == column_name:
            return column

    columns = [column for column in register.columns if column.casefold().startswith(column_name)]
    if len(columns) != 1:
        raise KeyError(f"Expected one column starting with '{column_name}', found {len(columns)}")

    return columns[0]


def get_value(register: pd.DataFrame, row, column_name: str) -> str:
    return row[find_column(register, column_name)].strip()


def join_texts(texts: list[str]) -> str:
    # Several values are shown as a list, one "- value" per line
    texts = [text for text in texts if text and text != "---"]
    if len(texts) > 1:
        texts = [f"- {text}" for text in texts]

    return "\n".join(texts)


def get_rows(register: pd.DataFrame, row, fields: list[tuple[str, list[str]]]) -> list[tuple[str, str]]:
    # The (title, value) of the fields with a value
    rows = list()
    for title, columns in fields:
        value = join_texts([get_value(register, row, column) for column in columns])
        if value:
            rows.append((title, value))

    return rows


# ---------- Images ----------

def to_pixels(length: int) -> int:
    return round(Emu(length).inches * IMAGE_DPI)


def to_png(image: Image.Image) -> bytes:
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


@lru_cache
def make_solid_image(color: str) -> bytes:
    return to_png(Image.new("RGB", (8, 8), f"#{color}"))


@lru_cache
def make_card_cap(theme: Theme, is_top: bool) -> bytes:
    # The top or the bottom of a card: the rounded textured frame with the rounded card inside and the ring at the end
    # of the timeline, with transparent corners. It's drawn as the top, the bottom is the same picture upside down.
    width, height = to_pixels(CONTENT_WIDTH), to_pixels(FRAME_WIDTH + CARD_RADIUS)
    frame, ring = to_pixels(FRAME_WIDTH), to_pixels(TIMELINE_RING_SIZE)

    # The frame: the top (or the bottom, upside down) of the texture, as wide as the card. The sides of the frame
    # below this picture are a single color, so below the top of the frame the texture fades into that color.
    texture = Image.open(theme.frame_texture_file_path).convert("RGB")
    texture = texture.resize((width, round(texture.height * width / texture.width)), Image.LANCZOS)
    if is_top:
        texture = texture.crop((0, 0, width, height))
    else:
        texture = texture.crop((0, texture.height - height, width, texture.height)).transpose(Image.FLIP_TOP_BOTTOM)
    fade = Image.new("L", (width, height), 0)
    fade_draw = ImageDraw.Draw(fade)
    for y in range(frame, height):
        fade_draw.line((0, y, width - 1, y), fill=round(255 * (y - frame) / (height - 1 - frame)))
    image =Image.composite(Image.new("RGB", (width, height), f"#{theme.frame_color}"), texture, fade).convert("RGBA")

    # The rounded shapes and the ring are drawn 4 times bigger and scaled down, so the curves are smooth
    scale = 4
    shape = Image.new("L", (width * scale, height * 4 * scale), 0)
    ImageDraw.Draw(shape).rounded_rectangle((0, 0, shape.width - 1, shape.height - 1),
                                            radius=to_pixels(FRAME_RADIUS) * scale, fill=255)
    image.putalpha(shape.crop((0, 0, width * scale, height * scale)).resize((width, height), Image.LANCZOS))

    card = Image.new("RGBA", (width * scale, height * 4 * scale), (0, 0, 0, 0))
    draw = ImageDraw.Draw(card)
    draw.rounded_rectangle((frame * scale, frame * scale, (width - frame) * scale - 1, card.height - 1),
                           radius=to_pixels(CARD_RADIUS) * scale, fill=f"#{theme.card_color}")
    # The ring is above the dotted line, which starts at the bottom of the picture
    ring_x = (frame + to_pixels(TIMELINE_WIDTH) // 2 - ring // 2) * scale
    ring_y = (height - ring - to_pixels(Cm(0.05))) * scale
    draw.ellipse((ring_x, ring_y, ring_x + ring * scale - 1, ring_y + ring * scale - 1),
                 outline=f"#{CARD_TEXT_COLOR}", width=round(ring * scale * 0.18))
    image.alpha_composite(card.crop((0, 0, width * scale, height * scale)).resize((width, height), Image.LANCZOS))

    return to_png(image if is_top else image.transpose(Image.FLIP_TOP_BOTTOM))


# ---------- Word XML ----------
# Word opens a file only if the child elements of a property are in the order of the OOXML schema,
# so the new elements are inserted before the ones that come after them.

TABLE_PROPERTIES_ORDER = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                          "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
                          "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
                          "w:tblPrChange"]
CELL_PROPERTIES_ORDER = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
                         "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                         "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange"]


def set_property(properties, element, order: list[str]):
    tag = element.tag.split("}")[1]
    for old in properties.findall(qn(f"w:{tag}")):
        properties.remove(old)
    properties.insert_element_before(element, *order[order.index(f"w:{tag}") + 1:])


def make_element(tag: str, attributes: dict[str, str], children: list = None):
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(qn(key), str(value))
    for child in children or []:
        element.append(child)
    return element


def make_margins(tag: str, top: int, left: int, bottom: int, right: int):
    # The margins are in twips (1/20 pt)
    return make_element(tag, {}, [make_element(f"w:{side}", {"w:w": size, "w:type": "dxa"})
                                  for side, size in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]])


def setup_table(table, column_widths: list[int]):
    # A fixed layout without borders
    table.autofit = False
    for column, width in zip(table.columns, column_widths):
        column.width = width  # some editors use only the column widths, others only the cell widths
    borders = make_element("w:tblBorders", {}, [make_element(f"w:{side}", {"w:val": "nil"}) for side in
                                                ["top", "left", "bottom", "right", "insideH", "insideV"]])
    set_property(table._tbl.tblPr, borders, TABLE_PROPERTIES_ORDER)
    # No default cell margins, so all the rows start at the page margin. The cells set their own top and bottom.
    set_property(table._tbl.tblPr, make_margins("w:tblCellMar", 0, 0, 0, 0), TABLE_PROPERTIES_ORDER)


def add_table_row(table, column_widths: list[int]):
    table_row = table.add_row()
    for cell, width in zip(table_row.cells, column_widths):
        cell.width = width
    return table_row


def set_cell_background(cell, color: str):
    shading = make_element("w:shd", {"w:val": "clear", "w:color": "auto", "w:fill": color})
    set_property(cell._tc.get_or_add_tcPr(), shading, CELL_PROPERTIES_ORDER)


def set_cell_margins(cell, top: int, left: int, bottom: int, right: int):
    set_property(cell._tc.get_or_add_tcPr(), make_margins("w:tcMar", top, left, bottom, right), CELL_PROPERTIES_ORDER)


def set_cell_padding(cell, top: int, left: int, bottom: int, right: int):
    # Some editors move a whole table row by the left margin of its first cell, which would misalign the rounded
    # corners of the cards. So only the top and bottom are cell margins, the sides are paragraph indents.
    set_cell_margins(cell, Emu(top).twips, 0, Emu(bottom).twips, 0)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.left_indent = left + (paragraph.paragraph_format.left_indent or 0)
        paragraph.paragraph_format.right_indent = right


def set_cell_borders(cell, borders: dict[str, tuple[str, str, int]]):
    # {side: (style, color, size in 1/8 pt)}
    elements = [make_element(f"w:{side}", {"w:val": style, "w:sz": size, "w:color": color})
                for side in ["top", "left", "bottom", "right"] if side in borders
                for style, color, size in [borders[side]]]
    set_property(cell._tc.get_or_add_tcPr(), make_element("w:tcBorders", {}, elements), CELL_PROPERTIES_ORDER)


def set_row_height(table_row, height: int):
    # An exact height, the content that doesn't fit is cut
    table_row._tr.get_or_add_trPr().append(make_element("w:trHeight", {"w:val": Emu(height).twips,
                                                                       "w:hRule": "exact"}))


def add_run(paragraph, text: str, size: float, color: str, bold: bool = False):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    return run


def add_page_number(paragraph, size: float, color: str):
    # A PAGE field. Every part of it is a separate run with the same look, otherwise the number gets the default look.
    for field_type in ["begin", "instruction", "separate", "number", "end"]:
        run = add_run(paragraph, "1" if field_type == "number" else "", size, color)
        if field_type == "instruction":
            element = make_element("w:instrText", {"xml:space": "preserve"})
            element.text = " PAGE "
            run._r.append(element)
        elif field_type != "number":
            run._r.append(make_element("w:fldChar", {"w:fldCharType": field_type}))


def add_inline_picture(run, image, width: int = None, height: int = None):
    # Without the distances to the text some editors put 1/8 inch of space around the picture
    inline = run.add_picture(image, width=width, height=height)._inline
    for side in ["distT", "distB", "distL", "distR"]:
        inline.set(side, "0")


def add_floating_picture(paragraph, image: bytes, x: int, y: int, width: int, height: int, z_order: int):
    # python-docx adds only inline pictures, so the picture is turned into a floating one behind the text,
    # placed relative to the page
    inline = paragraph.add_run().add_picture(io.BytesIO(image), width=width, height=height)._inline
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{z_order}" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{width}" cy="{height}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f'</wp:anchor>')
    for child in list(inline):
        if child.tag in [qn("wp:docPr"), qn("wp:cNvGraphicFramePr"), qn("a:graphic")]:
            anchor.append(child)
    inline.getparent().replace(inline, anchor)


# ---------- Profile ----------

def setup_styles(doc: docx.Document):
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor.from_string(CARD_TEXT_COLOR)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.15

    style = doc.styles.add_style(FIELD_TITLE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(9)
    style.font.bold = True
    style.font.all_caps = True  # only the look, the text keeps its case
    style.paragraph_format.space_before = Pt(11)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.keep_with_next = True

    style = doc.styles.add_style(FIELD_VALUE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(10.5)


def setup_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    # Header: the dark background, a big rooster in the corner (partly outside of the page) and the logo,
    # on every page of the section
    section.header.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_floating_picture(header, make_solid_image(PAGE_COLOR), 0, 0, section.page_width, section.page_height, 0)
    with open(WATERMARK_FILE_PATH, "rb") as fstream:
        watermark = fstream.read()
    watermark_size = Image.open(io.BytesIO(watermark)).size
    watermark_height = Emu(int(WATERMARK_WIDTH * watermark_size[1] / watermark_size[0]))
    add_floating_picture(header, watermark, section.page_width - Cm(8.5), section.page_height - Cm(14.5),
                         WATERMARK_WIDTH, watermark_height, 1)
    add_inline_picture(header.add_run(), LOGO_FILE_PATH, height=Cm(1.1))

    # Footer: the page number on every page of the section
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer, 9, SUBTITLE_COLOR)


def add_title(doc: docx.Document, title: str, subtitle: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    add_run(paragraph, title, 30, PAGE_TEXT_COLOR, bold=True)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    add_run(paragraph, "→ ", 12, ACCENT_COLOR, bold=True)
    add_run(paragraph, subtitle, 12, SUBTITLE_COLOR)


def add_card_cap_row(table, column_widths: list[int], theme: Theme, is_top: bool):
    # The rounded top or bottom of the frame and the card, over all the columns
    height = FRAME_WIDTH + CARD_RADIUS
    table_row = add_table_row(table, column_widths)
    set_row_height(table_row, height)
    cell = table_row.cells[0].merge(table_row.cells[-1])
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.keep_with_next = is_top
    run = paragraph.add_run()
    run.font.size = Pt(1)  # so the line isn't higher than the picture
    add_inline_picture(run, io.BytesIO(make_card_cap(theme, is_top)), CONTENT_WIDTH, height)


def add_card(doc: docx.Document, number: int, title: str, rows: list[tuple[str, str]], theme: Theme):
    # The title: a big number and the name of the section
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(22)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True
    add_run(heading, f"{number:02d}", 22, ACCENT_COLOR, bold=True)
    add_run(heading, f"  {title}", 15, PAGE_TEXT_COLOR, bold=True)

    # The card in its frame: a table with rounded top and bottom, and a single row between them, so no lines show
    # between the rows in some editors. On the left of the card there's a dotted timeline with a ring at both ends
    # (the rings are in the pictures of the top and the bottom), on the right a title and a value for every field.
    # The columns: the frame, the left half of the timeline with the dotted line on its right side, the text, the frame.
    half_timeline = Emu(int(TIMELINE_WIDTH / 2))
    column_widths = [FRAME_WIDTH, half_timeline, TIMELINE_WIDTH - half_timeline + TEXT_WIDTH, FRAME_WIDTH]
    table = doc.add_table(rows=0, cols=len(column_widths))
    setup_table(table, column_widths)
    add_card_cap_row(table, column_widths, theme, is_top=True)

    frame_left, timeline_cell, text_cell, frame_right = add_table_row(table, column_widths).cells
    for cell in [frame_left, frame_right]:
        set_cell_background(cell, theme.frame_color)
    for cell in [timeline_cell, text_cell]:
        set_cell_background(cell, theme.card_color)
    set_cell_borders(timeline_cell, {"right": ("dotted", CARD_TEXT_COLOR, TIMELINE_DOT_SIZE)})

    for row_i, (title, value) in enumerate(rows):
        paragraph = text_cell.paragraphs[0] if row_i == 0 else text_cell.add_paragraph()
        paragraph.style = doc.styles[FIELD_TITLE_STYLE]
        if row_i == 0:
            paragraph.paragraph_format.space_before = Pt(0)
        paragraph.add_run(title)

        # Every line of the value is a separate paragraph, the "- " lines have a hanging indent
        for line in value.split("\n"):
            paragraph = text_cell.add_paragraph(line, doc.styles[FIELD_VALUE_STYLE])
            if line.startswith("- "):
                paragraph.paragraph_format.left_indent = Cm(0.35)
                paragraph.paragraph_format.first_line_indent = Cm(-0.35)
                paragraph.paragraph_format.space_after = Pt(2)

    # The last line stays on the same page as the bottom of the card
    text_cell.paragraphs[-1].paragraph_format.keep_with_next = True
    set_cell_padding(text_cell, Pt(3), half_timeline + Cm(0.35), Pt(6), Cm(0.4))

    add_card_cap_row(table, column_widths, theme, is_top=False)


def create_doc(register: pd.DataFrame, row, file_path: str):
    doc = docx.Document()
    setup_styles(doc)
    setup_section(doc.sections[0])
    add_title(doc, PROFILE_TITLE, PROFILE_SUBTITLE)

    number = 1
    for section_title, fields in PROFILE_SECTIONS:
        rows = get_rows(register, row, fields)
        if rows:
            add_card(doc, number, section_title, rows, WARM_THEME)
            number += 1

    doc.save(file_path)


def create_docs():
    if not os.path.exists(OUTPUT_DIRECTORY):
        os.mkdir(OUTPUT_DIRECTORY)

    register = read_register(REGISTER_FILE_PATH)
    students = list()
    for _, row in register.iterrows():
        name = f"{get_value(register, row, STUDENT_FIRST_NAME)} {get_value(register, row, STUDENT_LAST_NAME)}".strip()
        if not name:
            continue

        status = get_value(register, row, STUDENT_STATUS)
        if status.casefold() not in PROFILE_STATUSES:
            # print(f"skip student {name}: {status or 'no status'}")
            continue

        students.append((name, row))

    # Students with equal names get a number suffix, so their files don't overwrite each other
    name_counts = dict()
    for name, _ in students:
        name_counts[name] = name_counts.get(name, 0) + 1

    name_counters = dict()
    for name, row in students:
        file_name = name
        if name_counts[name] > 1:
            name_counters[name] = name_counters.get(name, 0) + 1
            file_name = f"{name}_{name_counters[name]}"

        file_path = f"{OUTPUT_DIRECTORY}/{file_name}.docx"
        create_doc(register, row, file_path)
        print(f"created {file_path}")


if __name__ == "__main__":
    create_docs()
